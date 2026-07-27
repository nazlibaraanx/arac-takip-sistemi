import os
import tempfile
import unittest
from datetime import datetime, timedelta

from openpyxl import load_workbook

from excel_aktar import ExcelAktar
from sefer_yonu import BASKIL_MALATYA, MALATYA_BASKIL
from ucret_hesapla import UcretHesapla
from veritabani import Veritabani


class VeritabaniVeRaporTestleri(unittest.TestCase):

    def setUp(self):
        self.gecici_klasor = tempfile.TemporaryDirectory()
        self.veritabani_yolu = os.path.join(
            self.gecici_klasor.name,
            "test_feribot.db"
        )
        self.vt = Veritabani(self.veritabani_yolu)

    def tearDown(self):
        self.vt.baglantiyi_kapat()
        self.gecici_klasor.cleanup()

    def test_ucret_uzunluk_araligina_gore_secilir(self):
        self.assertEqual(self.vt.ucret_getir("Kamyon", 7.5), 550)
        self.assertEqual(self.vt.ucret_getir("Kamyon", 8.5), 800)
        self.assertEqual(self.vt.ucret_getir("Kamyon", 9.5), 1175)
        self.assertEqual(self.vt.ucret_getir("Kamyon", 6.5), 0)

    def test_uzunluk_yokken_orta_tarife_secilir(self):
        self.assertEqual(self.vt.ucret_getir("Kamyon"), 800)

    def test_otobus_toplu_tasima_tarifesine_eslenir(self):
        hesaplayici = UcretHesapla(self.vt)

        self.assertEqual(hesaplayici.ucret_hesapla("Otobüs"), 280)
        self.assertEqual(hesaplayici.ucret_hesapla("Servis"), 280)
        self.assertEqual(hesaplayici.ucret_hesapla("Minibüs"), 280)
        self.assertEqual(hesaplayici.ucret_hesapla("Toplu Taşıma"), 280)

    def test_panel_ozeti_alti_alandan_olusur(self):
        self.vt.gecis_kaydet(
            takip_id=42,
            plaka="34ABC123",
            arac_turu="Otomobil",
            genislik_px=180,
            ucret=240
        )

        kayit = self.vt.gecis_ozetlerini_getir()[0]

        self.assertEqual(len(kayit), 6)
        self.assertEqual(kayit[0], 1)
        self.assertEqual(kayit[3:], ("34ABC123", "Otomobil", 240.0))
        self.assertEqual(self.vt.plaka_ozetlerini_ara("ABC"), [kayit])

    def test_gecisler_tarih_araligina_gore_filtrelenir(self):
        self.vt.gecis_kaydet(1, "34ABC123", "Otomobil", 180, 240)
        self.vt.gecis_kaydet(2, "44DEF456", "Kamyon", 250, 800)
        bugun = datetime.now()
        eski_tarih = (bugun - timedelta(days=10)).strftime("%d.%m.%Y")
        self.vt.imlec.execute(
            "UPDATE gecisler SET tarih = ? WHERE takip_id = ?",
            (eski_tarih, 1)
        )
        self.vt.baglanti.commit()

        bugunun_kayitlari = self.vt.gecis_ozetlerini_filtrele(bugun, bugun)
        self.assertEqual(len(bugunun_kayitlari), 1)
        self.assertEqual(bugunun_kayitlari[0][3], "44DEF456")

        plaka_sonucu = self.vt.gecis_ozetlerini_filtrele(
            bugun - timedelta(days=20),
            bugun,
            plaka="ABC"
        )
        self.assertEqual(len(plaka_sonucu), 1)
        self.assertEqual(plaka_sonucu[0][3], "34ABC123")

    def test_gecisleri_sil_tarifeleri_korur_ve_idyi_sifirlar(self):
        self.vt.gecis_kaydet(1, "34ABC123", "Otomobil", 180, 240)
        gunluk_ozet = self.vt.gunluk_ozetleri_getir()
        simdi = datetime.now()
        grafik_ozeti = self.vt.aylik_grafik_ozeti_getir(
            simdi.year,
            simdi.month
        )

        self.vt.tum_gecisleri_sil()

        self.assertEqual(self.vt.toplam_arac_sayisi(), 0)
        self.assertEqual(self.vt.toplam_gelir(), 0)
        self.assertEqual(self.vt.gecis_ozetlerini_getir(), [])
        self.assertGreater(len(self.vt.ucretleri_getir()), 0)
        self.assertEqual(self.vt.gunluk_ozetleri_getir(), gunluk_ozet)
        self.assertEqual(
            self.vt.aylik_grafik_ozeti_getir(simdi.year, simdi.month),
            grafik_ozeti
        )

        self.vt.gecis_kaydet(2, "34DEF456", "Otomobil", 170, 240)
        self.assertEqual(self.vt.gecis_ozetlerini_getir()[0][0], 1)

    def test_aylik_grafik_ozeti_gunleri_ve_arac_turlerini_dondurur(self):
        self.vt.gecis_kaydet(1, "A", "Otomobil", 180, 240)
        self.vt.gecis_kaydet(2, "B", "Toplu Taşıma", 210, 280)
        simdi = datetime.now()
        bugun = simdi.strftime("%d.%m.%Y")

        ozet = self.vt.aylik_grafik_ozeti_getir(
            simdi.year,
            simdi.month
        )

        self.assertEqual(
            ozet["gunluk_ozetler"],
            [(bugun, 2, 0, 520.0)]
        )
        self.assertEqual(
            ozet["arac_turleri"],
            [("Otomobil", 1), ("Toplu Taşıma", 1)]
        )
        self.assertIn(
            (simdi.year, simdi.month),
            self.vt.rapor_aylarini_getir()
        )

    def test_grafik_ozeti_secili_tarih_araligina_gore_filtrelenir(self):
        self.vt.gecis_kaydet(1, "A", "Otomobil", 180, 240)
        bugun = datetime.now()
        eski_gun = bugun - timedelta(days=10)
        eski_tarih = eski_gun.strftime("%d.%m.%Y")
        self.vt.imlec.execute(
            """
            INSERT INTO gunluk_ozetler
                (tarih, sefer_yonu, toplam_arac, toplam_yaya, toplam_hasilat)
            VALUES (?, ?, ?, ?, ?)
            """,
            (eski_tarih, MALATYA_BASKIL, 4, 2, 960)
        )
        self.vt.imlec.execute(
            """
            INSERT INTO gunluk_arac_turleri
                (tarih, sefer_yonu, arac_turu, toplam_arac, toplam_hasilat)
            VALUES (?, ?, ?, ?, ?)
            """,
            (eski_tarih, MALATYA_BASKIL, "Kamyon", 4, 960)
        )
        self.vt.baglanti.commit()

        ozet = self.vt.tarih_araligi_grafik_ozeti_getir(bugun, bugun)
        self.assertEqual(len(ozet["gunluk_ozetler"]), 1)
        self.assertEqual(ozet["gunluk_ozetler"][0][1], 1)
        self.assertEqual(ozet["arac_turleri"], [("Otomobil", 1)])

    def test_excel_basliklari_ve_satiri_uyumludur(self):
        kayitlar = [
            (1, "21.07.2026", "12:30:00", "34ABC123", "car", 240),
            (2, "21.07.2026", "13:30:00", "34DEF456", "truck", 800),
            (3, "22.07.2026", "09:00:00", "34GHI789", "bus", 280)
        ]

        dosya_yolu = ExcelAktar.aktar(
            kayitlar,
            hedef_klasor=self.gecici_klasor.name
        )

        calisma_kitabi = load_workbook(dosya_yolu, read_only=True)
        sayfa = calisma_kitabi.active
        satirlar = list(sayfa.iter_rows(values_only=True))
        calisma_kitabi.close()

        self.assertEqual(
            satirlar[0][0],
            "Malatya - Baskil · Temmuz 2026 Geçiş Özeti"
        )
        self.assertEqual(len(satirlar[7]), 4)
        self.assertEqual(
            satirlar[7],
            (
                "Tarih",
                "Toplam Taşınan Araç",
                "Yaya Yolcu",
                "Toplam Hasılat (TL)"
            )
        )
        self.assertEqual(satirlar[8][0].strftime("%d.%m.%Y"), "21.07.2026")
        self.assertEqual(satirlar[8][1:], (2, 0, 1040))
        self.assertEqual(satirlar[9][0].strftime("%d.%m.%Y"), "22.07.2026")
        self.assertEqual(satirlar[9][1:], (1, 0, 280))
        self.assertEqual(satirlar[10][0], "AY TOPLAMI")
        self.assertEqual(satirlar[10][1], "=SUM(B9:B10)")
        self.assertEqual(satirlar[10][2], "=SUM(C9:C10)")
        self.assertEqual(satirlar[10][3], "=SUM(D9:D10)")
        self.assertEqual(satirlar[4][0], "=B11")
        self.assertEqual(satirlar[4][1], "=C11")
        self.assertEqual(satirlar[4][2], "=D11")

    def test_gunluk_ozet_ayni_tarihteki_kayitlari_birlestirir(self):
        kayitlar = [
            (1, "21.07.2026", "10:00:00", "A", "car", 240),
            (2, "21.07.2026", "11:00:00", "B", "truck", 800)
        ]

        ozet = ExcelAktar.gunluk_ozet_hazirla(kayitlar)

        self.assertEqual(len(ozet), 1)
        self.assertEqual(ozet[0][0], "21.07.2026")
        self.assertEqual(ozet[0][1:], (2, 0, 1040))

    def test_ay_icindeki_guncellemeler_ayni_excel_dosyasini_kullanir(self):
        ilk_kayitlar = [
            (1, "21.07.2026", "10:00:00", "A", "car", 240)
        ]
        tum_kayitlar = ilk_kayitlar + [
            (2, "22.07.2026", "11:00:00", "B", "truck", 800)
        ]

        ilk_dosya = ExcelAktar.aktar(
            ilk_kayitlar,
            hedef_klasor=self.gecici_klasor.name
        )
        ikinci_dosya = ExcelAktar.aktar(
            tum_kayitlar,
            hedef_klasor=self.gecici_klasor.name
        )

        self.assertEqual(ilk_dosya, ikinci_dosya)
        self.assertTrue(os.path.exists(ikinci_dosya))

        calisma_kitabi = load_workbook(ikinci_dosya, read_only=True)
        sayfa = calisma_kitabi.active
        tarihler = [
            satir[0].strftime("%d.%m.%Y")
            for satir in list(sayfa.iter_rows(values_only=True))[8:10]
        ]
        calisma_kitabi.close()

        self.assertEqual(tarihler, ["21.07.2026", "22.07.2026"])

    def test_farkli_aylar_icin_ayri_excel_dosyalari_olusturur(self):
        ozetler = [
            ("31.07.2026", 4, 1040),
            ("01.08.2026", 2, 520)
        ]

        dosyalar = ExcelAktar.aylik_raporlari_guncelle(
            ozetler,
            hedef_klasor=self.gecici_klasor.name,
            referans_tarihi="01.08.2026"
        )

        beklenen_dosyalar = {
            "Feribot_Aylik_Gecis_Raporu_Malatya_Baskil_2026_07_Temmuz.xlsx",
            "Feribot_Aylik_Gecis_Raporu_Malatya_Baskil_2026_08_Ağustos.xlsx"
        }
        self.assertEqual(
            {os.path.basename(dosya) for dosya in dosyalar},
            beklenen_dosyalar
        )

        temmuz_yolu = os.path.join(
            self.gecici_klasor.name,
            "Feribot_Aylik_Gecis_Raporu_Malatya_Baskil_2026_07_Temmuz.xlsx"
        )
        agustos_yolu = os.path.join(
            self.gecici_klasor.name,
            "Feribot_Aylik_Gecis_Raporu_Malatya_Baskil_2026_08_Ağustos.xlsx"
        )
        self.assertTrue(os.path.exists(temmuz_yolu))
        self.assertTrue(os.path.exists(agustos_yolu))

        temmuz_kitabi = load_workbook(temmuz_yolu, read_only=True)
        agustos_kitabi = load_workbook(agustos_yolu, read_only=True)
        self.assertEqual(temmuz_kitabi.active.title, "Temmuz 2026")
        self.assertEqual(agustos_kitabi.active.title, "Ağustos 2026")
        temmuz_kitabi.close()
        agustos_kitabi.close()

    def test_dosya_adi_yil_ay_ve_turkce_ay_adini_icerir(self):
        self.assertEqual(
            ExcelAktar.dosya_adi_olustur(2026, 7),
            "Feribot_Aylik_Gecis_Raporu_Malatya_Baskil_2026_07_Temmuz.xlsx"
        )

    def test_iki_sefer_yonunun_verileri_ve_excel_dosyalari_ayridir(self):
        bugun = datetime.now().strftime("%d.%m.%Y")
        self.vt.gecis_kaydet(
            1, "44MB001", "Otomobil", 180, 240, MALATYA_BASKIL
        )
        self.vt.gecis_kaydet(
            2, "44BM001", "Toplu Taşıma", 220, 280, BASKIL_MALATYA
        )
        self.vt.yaya_kaydet(101, MALATYA_BASKIL)
        self.vt.yaya_kaydet(102, BASKIL_MALATYA)
        self.vt.yaya_kaydet(103, BASKIL_MALATYA)

        self.assertEqual(
            self.vt.gunluk_ozetleri_getir(MALATYA_BASKIL),
            [(bugun, 1, 1, 240.0)]
        )
        self.assertEqual(
            self.vt.gunluk_ozetleri_getir(BASKIL_MALATYA),
            [(bugun, 1, 2, 280.0)]
        )
        self.assertEqual(
            self.vt.gecis_ozetlerini_getir(MALATYA_BASKIL)[0][3],
            "44MB001"
        )
        self.assertEqual(
            self.vt.gecis_ozetlerini_getir(BASKIL_MALATYA)[0][3],
            "44BM001"
        )

        malatya_dosyasi = ExcelAktar.gunluk_raporu_guncelle(
            self.vt.gunluk_ozetleri_getir(MALATYA_BASKIL),
            hedef_klasor=self.gecici_klasor.name,
            sefer_yonu=MALATYA_BASKIL
        )
        baskil_dosyasi = ExcelAktar.gunluk_raporu_guncelle(
            self.vt.gunluk_ozetleri_getir(BASKIL_MALATYA),
            hedef_klasor=self.gecici_klasor.name,
            sefer_yonu=BASKIL_MALATYA
        )
        self.assertNotEqual(malatya_dosyasi, baskil_dosyasi)
        self.assertTrue(os.path.exists(malatya_dosyasi))
        self.assertTrue(os.path.exists(baskil_dosyasi))

    def test_rapor_merkezi_butun_yon_ve_aylari_tek_dosyada_toplar(self):
        eski_dosya = os.path.join(
            self.gecici_klasor.name,
            "Feribot_Aylik_Gecis_Raporu_Malatya_Baskil_2026_07_Temmuz.xlsx",
        )
        with open(eski_dosya, "wb") as dosya:
            dosya.write(b"eski")

        rapor_yolu = ExcelAktar.rapor_merkezini_guncelle(
            {
                MALATYA_BASKIL: [
                    ("24.07.2026", 5, 1, 1200),
                    ("01.08.2026", 3, 0, 720),
                ],
                BASKIL_MALATYA: [
                    ("24.07.2026", 4, 2, 960),
                ],
            },
            sefer_ozetleri=[
                (
                    1, "SF-20260724-001", MALATYA_BASKIL,
                    "24.07.2026", "08:00:00",
                    "24.07.2026", "08:30:00", "Tamamlandı",
                    5, 1, 1200.0,
                )
            ],
            hedef_klasor=self.gecici_klasor.name,
        )

        self.assertEqual(
            os.path.basename(rapor_yolu),
            "Feribot_Rapor_Merkezi.xlsx",
        )
        self.assertTrue(os.path.exists(rapor_yolu))
        self.assertFalse(os.path.exists(eski_dosya))

        calisma_kitabi = load_workbook(rapor_yolu, data_only=False)
        self.assertEqual(
            calisma_kitabi.sheetnames,
            [
                "Rapor Merkezi",
                "Aylık Genel",
                "Yön Bazlı",
                "Sefer Bazlı",
                "Grafikler",
            ],
        )
        sayfa = calisma_kitabi["Rapor Merkezi"]
        self.assertEqual(sayfa["A1"].value, "FERİBOT RAPOR MERKEZİ")
        self.assertEqual(
            [sayfa.cell(7, sutun).value for sutun in range(1, 7)],
            [
                "Sefer Yönü",
                "Dönem",
                "Tarih",
                "Toplam Araç",
                "Yaya Yolcu",
                "Toplam Hasılat (TL)",
            ],
        )
        self.assertEqual(sayfa["A8"].value, MALATYA_BASKIL)
        self.assertEqual(sayfa["B8"].value, "Ağustos 2026")
        self.assertEqual(sayfa["A9"].value, MALATYA_BASKIL)
        self.assertEqual(sayfa["A10"].value, BASKIL_MALATYA)
        self.assertEqual(sayfa["B4"].value, "=SUBTOTAL(109,D8:D10)")
        self.assertEqual(sayfa["D4"].value, "=SUBTOTAL(109,E8:E10)")
        self.assertEqual(sayfa["F4"].value, "=SUBTOTAL(109,F8:F10)")
        self.assertEqual(sayfa.auto_filter.ref, "A7:F10")
        genel = calisma_kitabi["Aylık Genel"]
        self.assertEqual(
            [genel.cell(5, sutun).value for sutun in range(1, 5)],
            ["Ağustos 2026", 3, 0, 720.0],
        )
        self.assertEqual(
            [genel.cell(6, sutun).value for sutun in range(1, 5)],
            ["Temmuz 2026", 9, 3, 2160.0],
        )
        sefer = calisma_kitabi["Sefer Bazlı"]
        self.assertEqual(sefer["A5"].value, "SF-20260724-001")
        self.assertEqual(sefer["H5"].value, 5)
        self.assertEqual(len(calisma_kitabi["Grafikler"]._charts), 2)
        calisma_kitabi.close()

    def test_gecis_kaydi_kalici_gunluk_toplami_artirir(self):
        bugun = datetime.now().strftime("%d.%m.%Y")

        self.vt.gecis_kaydet(1, "A", "Otomobil", 180, 240)
        self.vt.gecis_kaydet(2, "B", "Toplu Taşıma", 200, 280)

        self.assertEqual(
            self.vt.gunluk_ozetleri_getir(),
            [(bugun, 2, 0, 520.0)]
        )

    def test_yukleme_araci_yayayi_ve_hasilati_tek_seferde_toplar(self):
        sefer_id = self.vt.sefer_baslat(MALATYA_BASKIL)
        self.vt.gecis_kaydet(
            1, "44SF001", "Otomobil", 180, 240,
            MALATYA_BASKIL, sefer_id
        )
        self.vt.gecis_kaydet(
            2, "44SF002", "Toplu Taşıma", 220, 280,
            MALATYA_BASKIL, sefer_id
        )
        self.assertTrue(
            self.vt.yaya_kaydet(90, MALATYA_BASKIL, sefer_id)
        )
        self.assertTrue(self.vt.sefer_bitir(sefer_id))

        sefer = self.vt.sefer_ozetlerini_getir()[0]
        self.assertEqual(sefer[0], sefer_id)
        self.assertTrue(sefer[1].startswith("SF-"))
        self.assertEqual(sefer[2], MALATYA_BASKIL)
        self.assertEqual(sefer[7], "Tamamlandı")
        self.assertEqual(sefer[8:], (2, 1, 520.0))

    def test_yaya_ayni_takip_idsiyle_yalnizca_bir_kez_sayilir(self):
        bugun = datetime.now().strftime("%d.%m.%Y")

        self.assertTrue(self.vt.yaya_kaydet(15))
        self.assertFalse(self.vt.yaya_kaydet(15))
        self.assertTrue(self.vt.yaya_kaydet(16))

        self.assertEqual(self.vt.toplam_yaya_sayisi(), 2)
        self.assertEqual(
            self.vt.gunluk_ozetleri_getir(),
            [(bugun, 0, 2, 0.0)]
        )

    def test_yaya_takip_idsi_yeni_oturumda_yeniden_kullanilabilir(self):
        self.assertTrue(self.vt.yaya_kaydet(15))
        ikinci_oturum = Veritabani(self.veritabani_yolu)
        try:
            self.assertTrue(ikinci_oturum.yaya_kaydet(15))
            self.assertEqual(ikinci_oturum.toplam_yaya_sayisi(), 2)
        finally:
            ikinci_oturum.baglantiyi_kapat()

    def test_rapor_verilerini_sil_tarifeleri_korur(self):
        self.vt.gecis_kaydet(1, "A", "Otomobil", 180, 240)
        self.vt.yaya_kaydet(2)

        self.vt.tum_rapor_verilerini_sil()

        self.assertEqual(self.vt.gunluk_ozetleri_getir(), [])
        self.assertEqual(self.vt.toplam_arac_sayisi(), 0)
        self.assertEqual(self.vt.toplam_yaya_sayisi(), 0)
        self.assertGreater(len(self.vt.ucretleri_getir()), 0)

    def test_excel_arac_turlerini_turkcelestirir(self):
        eslemeler = {
            "car": "Otomobil",
            "truck": "Kamyon",
            "bus": "Toplu Taşıma",
            "motorcycle": "Motosiklet",
            "minibus": "Toplu Taşıma"
        }

        for ingilizce, turkce in eslemeler.items():
            with self.subTest(arac_turu=ingilizce):
                self.assertEqual(
                    ExcelAktar.arac_turunu_turkcelestir(ingilizce),
                    turkce
                )


if __name__ == "__main__":
    unittest.main()
