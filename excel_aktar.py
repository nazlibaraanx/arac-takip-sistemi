import os
from collections import defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from ayarlar import RAPOR_KLASORU
from sefer_yonu import (
    BASKIL_MALATYA,
    MALATYA_BASKIL,
    dosya_etiketi,
    sefer_yonunu_dogrula,
)


class ExcelAktar:

    AY_ADLARI = {
        1: "Ocak",
        2: "Şubat",
        3: "Mart",
        4: "Nisan",
        5: "Mayıs",
        6: "Haziran",
        7: "Temmuz",
        8: "Ağustos",
        9: "Eylül",
        10: "Ekim",
        11: "Kasım",
        12: "Aralık"
    }

    RAPOR_DOSYA_SABLONU = (
        "Feribot_Aylik_Gecis_Raporu_{sefer}_{yil}_{ay:02d}_{ay_adi}.xlsx"
    )
    RAPOR_MERKEZI_DOSYASI = "Feribot_Rapor_Merkezi.xlsx"

    ARAC_ISIMLERI = {
        "car": "Otomobil",
        "truck": "Kamyon",
        "bus": "Toplu Taşıma",
        "motorcycle": "Motosiklet",
        "minibus": "Toplu Taşıma",
        "bicycle": "Bisiklet",
        "train": "Tren",
        "Otobus": "Toplu Taşıma",
        "Otobüs": "Toplu Taşıma",
        "OtobÃ¼s": "Toplu Taşıma",
        "Servis": "Toplu Taşıma",
        "Minibüs": "Toplu Taşıma",
        "MinibÃ¼s": "Toplu Taşıma",
        "Tir": "Uzun Araç",
        "Tır": "Uzun Araç",
        "TÄ±r": "Uzun Araç"
    }

    @classmethod
    def arac_turunu_turkcelestir(cls, arac_turu):
        return cls.ARAC_ISIMLERI.get(arac_turu, arac_turu)

    @staticmethod
    def gunluk_ozet_hazirla(kayitlar):
        """Eski biçimdeki geçiş kayıtlarını tarih bazında özetler."""

        gunler = {}

        for kayit in kayitlar:
            if len(kayit) < 6:
                continue

            tarih = kayit[1]
            ucret = float(kayit[5] or 0)
            gun = gunler.setdefault(
                tarih,
                {"arac_sayisi": 0, "yaya_sayisi": 0, "toplam_hasilat": 0.0}
            )
            gun["arac_sayisi"] += 1
            gun["toplam_hasilat"] += ucret

        return [
            (
                tarih,
                gunler[tarih]["arac_sayisi"],
                gunler[tarih]["yaya_sayisi"],
                gunler[tarih]["toplam_hasilat"]
            )
            for tarih in sorted(
                gunler,
                key=lambda deger: ExcelAktar._tarih_coz(deger)
                or datetime.min
            )
        ]

    @staticmethod
    def _tarih_coz(tarih):
        if isinstance(tarih, datetime):
            return tarih
        if isinstance(tarih, date):
            return datetime.combine(tarih, datetime.min.time())

        try:
            return datetime.strptime(str(tarih), "%d.%m.%Y")
        except (TypeError, ValueError):
            return None

    @classmethod
    def dosya_adi_olustur(cls, yil, ay, sefer_yonu=MALATYA_BASKIL):
        """Sıralanabilir ve okunabilir profesyonel aylık dosya adını üretir."""

        return cls.RAPOR_DOSYA_SABLONU.format(
            yil=int(yil),
            ay=int(ay),
            ay_adi=cls.AY_ADLARI[int(ay)],
            sefer=dosya_etiketi(sefer_yonu),
        )

    @classmethod
    def _ozetleri_aylara_ayir(cls, gunluk_ozetler):
        aylik_ozetler = defaultdict(list)

        for ozet_satiri in gunluk_ozetler:
            if len(ozet_satiri) == 3:
                tarih, toplam_arac, toplam_hasilat = ozet_satiri
                toplam_yaya = 0
            else:
                tarih, toplam_arac, toplam_yaya, toplam_hasilat = ozet_satiri
            tarih_degeri = cls._tarih_coz(tarih)
            if tarih_degeri is None:
                continue

            aylik_ozetler[(tarih_degeri.year, tarih_degeri.month)].append(
                (
                    tarih_degeri.date(),
                    int(toplam_arac),
                    int(toplam_yaya),
                    float(toplam_hasilat)
                )
            )

        for satirlar in aylik_ozetler.values():
            satirlar.sort(key=lambda satir: satir[0])

        return aylik_ozetler

    @classmethod
    def _aylik_raporu_yaz(
        cls, yil, ay, satirlar, hedef_klasor, sefer_yonu=MALATYA_BASKIL
    ):
        sefer_yonu = sefer_yonunu_dogrula(sefer_yonu)
        ay_adi = cls.AY_ADLARI[ay]
        dosya_adi = cls.dosya_adi_olustur(yil, ay, sefer_yonu)
        dosya_yolu = os.path.join(hedef_klasor, dosya_adi)
        gecici_dosya = os.path.join(
            hedef_klasor,
            f".{os.getpid()}_{dosya_adi}.tmp.xlsx"
        )

        wb = Workbook()
        ws = wb.active
        ws.title = f"{ay_adi} {yil}"
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:D1")
        ws["A1"] = f"{sefer_yonu} · {ay_adi} {yil} Geçiş Özeti"
        ws["A1"].fill = PatternFill("solid", fgColor="16324F")
        ws["A1"].font = Font(
            name="Aptos Display", color="FFFFFF", bold=True, size=18
        )
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 38

        ws.merge_cells("A2:D2")
        ws["A2"] = (
            "Günlük araç, yaya yolcu ve hasılat kayıtlarının aylık özeti"
        )
        ws["A2"].font = Font(name="Aptos", color="64748B", size=10)
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 22

        kart_dolgulari = ("E8F4F1", "EEF4FA", "FFF4E8")
        kart_basliklari = ("TOPLAM ARAÇ", "YAYA YOLCU", "TOPLAM HASILAT")
        ws.merge_cells("C4:D4")
        ws.merge_cells("C5:D5")
        for hucre, metin, renk in zip(
            ("A4", "B4", "C4"), kart_basliklari, kart_dolgulari
        ):
            ws[hucre] = metin
            ws[hucre].fill = PatternFill("solid", fgColor=renk)
            ws[hucre].font = Font(
                name="Aptos", color="64748B", bold=True, size=9
            )
            ws[hucre].alignment = Alignment(
                horizontal="center", vertical="center"
            )
        for hucre, renk in zip(("A5", "B5", "C5"), kart_dolgulari):
            ws[hucre].fill = PatternFill("solid", fgColor=renk)
            ws[hucre].font = Font(
                name="Aptos Display", color="16324F", bold=True, size=16
            )
            ws[hucre].alignment = Alignment(
                horizontal="center", vertical="center"
            )
        ws["D4"].fill = PatternFill("solid", fgColor=kart_dolgulari[2])
        ws["D5"].fill = PatternFill("solid", fgColor=kart_dolgulari[2])
        ws.row_dimensions[4].height = 21
        ws.row_dimensions[5].height = 30

        ws.merge_cells("A7:D7")
        ws["A7"] = "GÜNLÜK GEÇİŞ KAYITLARI"
        ws["A7"].font = Font(
            name="Aptos", color="16324F", bold=True, size=10
        )
        ws["A7"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[7].height = 23

        basliklar = [
            "Tarih",
            "Toplam Taşınan Araç",
            "Yaya Yolcu",
            "Toplam Hasılat (TL)"
        ]
        for sutun, baslik in enumerate(basliklar, start=1):
            ws.cell(8, sutun, baslik)

        ince_cizgi = Side(style="thin", color="DDE5EC")
        kenarlik = Border(
            bottom=ince_cizgi
        )

        for hucre in ws[8]:
            hucre.fill = PatternFill("solid", fgColor="2F6F77")
            hucre.font = Font(
                name="Aptos", color="FFFFFF", bold=True, size=10
            )
            hucre.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[8].height = 26

        for sira, (tarih, toplam_arac, toplam_yaya, toplam_hasilat) in enumerate(
            satirlar,
            start=9
        ):
            ws.cell(sira, 1, tarih)
            ws.cell(sira, 2, toplam_arac)
            ws.cell(sira, 3, toplam_yaya)
            ws.cell(sira, 4, toplam_hasilat)
            dolgu_rengi = "F7FAFC" if sira % 2 else "FFFFFF"

            for sutun in range(1, 5):
                hucre = ws.cell(sira, sutun)
                hucre.fill = PatternFill("solid", fgColor=dolgu_rengi)
                hucre.border = kenarlik
                hucre.font = Font(name="Aptos", color="334155", size=10)
                hucre.alignment = Alignment(
                    horizontal="center" if sutun < 4 else "right",
                    vertical="center"
                )
            ws.row_dimensions[sira].height = 22

            ws.cell(sira, 1).number_format = "dd.mm.yyyy"
            ws.cell(sira, 2).number_format = "#,##0"
            ws.cell(sira, 3).number_format = "#,##0"
            ws.cell(sira, 4).number_format = '#,##0.00 "TL"'

        ilk_veri_satiri = 9
        son_veri_satiri = 8 + len(satirlar)
        toplam_satiri = son_veri_satiri + 1

        ws.cell(toplam_satiri, 1, "AY TOPLAMI")
        if satirlar:
            ws.cell(
                toplam_satiri,
                2,
                f"=SUM(B{ilk_veri_satiri}:B{son_veri_satiri})"
            )
            ws.cell(
                toplam_satiri,
                3,
                f"=SUM(C{ilk_veri_satiri}:C{son_veri_satiri})"
            )
            ws.cell(
                toplam_satiri,
                4,
                f"=SUM(D{ilk_veri_satiri}:D{son_veri_satiri})"
            )
        else:
            ws.cell(toplam_satiri, 2, 0)
            ws.cell(toplam_satiri, 3, 0)
            ws.cell(toplam_satiri, 4, 0.0)

        for hucre in ws[toplam_satiri]:
            hucre.fill = PatternFill("solid", fgColor="E8F4F1")
            hucre.font = Font(
                name="Aptos", color="16324F", bold=True, size=10
            )
            hucre.border = Border(
                top=Side(style="medium", color="2F6F77")
            )
            hucre.alignment = Alignment(
                horizontal="center" if hucre.column < 4 else "right"
            )
        ws.row_dimensions[toplam_satiri].height = 25

        ws.cell(toplam_satiri, 2).number_format = "#,##0"
        ws.cell(toplam_satiri, 3).number_format = "#,##0"
        ws.cell(toplam_satiri, 4).number_format = '#,##0.00 "TL"'

        ws["A5"] = f"=B{toplam_satiri}"
        ws["B5"] = f"=C{toplam_satiri}"
        ws["C5"] = f"=D{toplam_satiri}"
        ws["A5"].number_format = "#,##0"
        ws["B5"].number_format = "#,##0"
        ws["C5"].number_format = '#,##0.00 "TL"'

        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 26
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 26
        ws.freeze_panes = "A9"
        ws.auto_filter.ref = (
            f"A8:D{son_veri_satiri}" if satirlar else "A8:D8"
        )
        ws.print_title_rows = "1:8"
        ws.print_area = f"A1:D{toplam_satiri}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.sheet_properties.pageSetUpPr.autoPageBreaks = False

        try:
            wb.save(gecici_dosya)
            wb.close()
            os.replace(gecici_dosya, dosya_yolu)
        finally:
            if os.path.exists(gecici_dosya):
                os.remove(gecici_dosya)

        return dosya_yolu

    @classmethod
    def rapor_merkezini_guncelle(
        cls,
        yonlere_gore_ozetler,
        sefer_ozetleri=None,
        hedef_klasor=RAPOR_KLASORU,
        eski_dosyalari_temizle=True,
    ):
        """Bütün yön ve ayları filtrelenebilir tek Excel dosyasında toplar."""

        os.makedirs(hedef_klasor, exist_ok=True)
        dosya_yolu = os.path.join(
            hedef_klasor,
            cls.RAPOR_MERKEZI_DOSYASI,
        )
        gecici_dosya = os.path.join(
            hedef_klasor,
            f".{os.getpid()}_{cls.RAPOR_MERKEZI_DOSYASI}.tmp.xlsx",
        )

        satirlar = []
        for sefer_yonu in (MALATYA_BASKIL, BASKIL_MALATYA):
            for tarih, toplam_arac, toplam_yaya, toplam_hasilat in (
                yonlere_gore_ozetler.get(sefer_yonu, [])
            ):
                tarih_degeri = cls._tarih_coz(tarih)
                if tarih_degeri is None:
                    continue
                satirlar.append(
                    (
                        sefer_yonu,
                        f"{cls.AY_ADLARI[tarih_degeri.month]} "
                        f"{tarih_degeri.year}",
                        tarih_degeri.date(),
                        int(toplam_arac),
                        int(toplam_yaya),
                        float(toplam_hasilat),
                    )
                )
        satirlar.sort(key=lambda satir: (satir[2], satir[0]), reverse=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Rapor Merkezi"
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:F1")
        ws["A1"] = "FERİBOT RAPOR MERKEZİ"
        ws["A1"].fill = PatternFill("solid", fgColor="16324F")
        ws["A1"].font = Font(
            name="Aptos Display", color="FFFFFF", bold=True, size=18
        )
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 38

        ws.merge_cells("A2:F2")
        ws["A2"] = (
            "Sefer Yönü ve Dönem başlıklarındaki filtrelerden istediğiniz "
            "raporu seçin. Üst toplamlar seçime göre otomatik güncellenir."
        )
        ws["A2"].font = Font(name="Aptos", color="64748B", size=10)
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 24

        ilk_veri_satiri = 8
        son_veri_satiri = max(ilk_veri_satiri, ilk_veri_satiri + len(satirlar) - 1)
        ozetler = (
            ("A4", "GÖRÜNEN ARAÇ", "B4", f"=SUBTOTAL(109,D{ilk_veri_satiri}:D{son_veri_satiri})"),
            ("C4", "GÖRÜNEN YAYA", "D4", f"=SUBTOTAL(109,E{ilk_veri_satiri}:E{son_veri_satiri})"),
            ("E4", "GÖRÜNEN HASILAT", "F4", f"=SUBTOTAL(109,F{ilk_veri_satiri}:F{son_veri_satiri})"),
        )
        for etiket_hucresi, etiket, deger_hucresi, formul in ozetler:
            ws[etiket_hucresi] = etiket
            ws[etiket_hucresi].fill = PatternFill("solid", fgColor="E8F4F1")
            ws[etiket_hucresi].font = Font(
                name="Aptos", color="64748B", bold=True, size=9
            )
            ws[etiket_hucresi].alignment = Alignment(
                horizontal="right", vertical="center"
            )
            ws[deger_hucresi] = formul
            ws[deger_hucresi].fill = PatternFill("solid", fgColor="E8F4F1")
            ws[deger_hucresi].font = Font(
                name="Aptos Display", color="16324F", bold=True, size=14
            )
            ws[deger_hucresi].alignment = Alignment(
                horizontal="left", vertical="center"
            )
        ws["B4"].number_format = "#,##0"
        ws["D4"].number_format = "#,##0"
        ws["F4"].number_format = '#,##0.00 "TL"'
        ws.row_dimensions[4].height = 31

        ws.merge_cells("A6:F6")
        ws["A6"] = "TÜM GÜNLÜK KAYITLAR"
        ws["A6"].font = Font(
            name="Aptos", color="16324F", bold=True, size=10
        )
        ws["A6"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[6].height = 23

        basliklar = (
            "Sefer Yönü",
            "Dönem",
            "Tarih",
            "Toplam Araç",
            "Yaya Yolcu",
            "Toplam Hasılat (TL)",
        )
        for sutun, baslik in enumerate(basliklar, start=1):
            hucre = ws.cell(7, sutun, baslik)
            hucre.fill = PatternFill("solid", fgColor="2F6F77")
            hucre.font = Font(
                name="Aptos", color="FFFFFF", bold=True, size=10
            )
            hucre.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[7].height = 27

        ince_cizgi = Side(style="thin", color="DDE5EC")
        for sira, satir in enumerate(satirlar, start=ilk_veri_satiri):
            for sutun, deger in enumerate(satir, start=1):
                hucre = ws.cell(sira, sutun, deger)
                hucre.fill = PatternFill(
                    "solid",
                    fgColor="F7FAFC" if sira % 2 else "FFFFFF",
                )
                hucre.font = Font(name="Aptos", color="334155", size=10)
                hucre.border = Border(bottom=ince_cizgi)
                hucre.alignment = Alignment(
                    horizontal="left" if sutun == 1 else (
                        "right" if sutun >= 4 else "center"
                    ),
                    vertical="center",
                )
            ws.cell(sira, 3).number_format = "dd.mm.yyyy"
            ws.cell(sira, 4).number_format = "#,##0"
            ws.cell(sira, 5).number_format = "#,##0"
            ws.cell(sira, 6).number_format = '#,##0.00 "TL"'
            ws.row_dimensions[sira].height = 22

        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 17
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 23
        ws.freeze_panes = "A8"
        ws.auto_filter.ref = f"A7:F{son_veri_satiri}"
        ws.print_title_rows = "1:7"
        ws.print_area = f"A1:F{son_veri_satiri}"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.sheet_properties.pageSetUpPr.autoPageBreaks = False

        aylik_genel = defaultdict(lambda: [0, 0, 0.0])
        aylik_yon = defaultdict(lambda: [0, 0, 0.0])
        for yon, donem, tarih, arac, yaya, hasilat in satirlar:
            anahtar = (tarih.year, tarih.month)
            aylik_genel[anahtar][0] += arac
            aylik_genel[anahtar][1] += yaya
            aylik_genel[anahtar][2] += hasilat
            aylik_yon[(yon, tarih.year, tarih.month)][0] += arac
            aylik_yon[(yon, tarih.year, tarih.month)][1] += yaya
            aylik_yon[(yon, tarih.year, tarih.month)][2] += hasilat

        genel_ws = wb.create_sheet("Aylık Genel")
        cls._standart_rapor_sayfasi(
            genel_ws,
            "AYLIK GENEL RAPOR",
            "İki sefer yönünün birlikte aylık toplamları",
            ("Dönem", "Toplam Araç", "Yaya Yolcu", "Toplam Hasılat (TL)"),
        )
        genel_satirlar = []
        for (yil, ay), degerler in sorted(aylik_genel.items(), reverse=True):
            genel_satirlar.append(
                (f"{cls.AY_ADLARI[ay]} {yil}", *degerler)
            )
        cls._rapor_satirlarini_yaz(genel_ws, genel_satirlar)

        yon_ws = wb.create_sheet("Yön Bazlı")
        cls._standart_rapor_sayfasi(
            yon_ws,
            "YÖN BAZLI AYLIK RAPOR",
            "Malatya–Baskil ve Baskil–Malatya aylık sonuçları ayrı gösterilir",
            (
                "Sefer Yönü", "Dönem", "Toplam Araç",
                "Yaya Yolcu", "Toplam Hasılat (TL)",
            ),
        )
        yon_satirlari = []
        for (yon, yil, ay), degerler in sorted(
            aylik_yon.items(),
            key=lambda oge: (oge[0][1], oge[0][2], oge[0][0]),
            reverse=True,
        ):
            yon_satirlari.append(
                (yon, f"{cls.AY_ADLARI[ay]} {yil}", *degerler)
            )
        cls._rapor_satirlarini_yaz(yon_ws, yon_satirlari)

        sefer_ws = wb.create_sheet("Sefer Bazlı")
        cls._standart_rapor_sayfasi(
            sefer_ws,
            "SEFER BAZLI RAPOR",
            "Her Yüklemeyi Başlat–Bitir aralığı ayrı sefer olarak saklanır",
            (
                "Sefer No", "Sefer Yönü", "Başlangıç Tarihi",
                "Başlangıç", "Bitiş Tarihi", "Bitiş", "Durum",
                "Araç", "Yaya", "Hasılat (TL)",
            ),
        )
        sefer_satirlari = [
            (
                satir[1], satir[2], satir[3], satir[4],
                satir[5] or "", satir[6] or "", satir[7],
                satir[8], satir[9], satir[10],
            )
            for satir in (sefer_ozetleri or [])
        ]
        cls._rapor_satirlarini_yaz(sefer_ws, sefer_satirlari)

        grafik_ws = wb.create_sheet("Grafikler")
        cls._grafik_sayfasini_hazirla(
            grafik_ws, genel_satirlar, yon_satirlari
        )

        try:
            wb.save(gecici_dosya)
            wb.close()
            os.replace(gecici_dosya, dosya_yolu)
        finally:
            if os.path.exists(gecici_dosya):
                os.remove(gecici_dosya)

        if eski_dosyalari_temizle:
            cls.eski_aylik_raporlari_temizle(hedef_klasor)
        return dosya_yolu

    @classmethod
    def _standart_rapor_sayfasi(cls, ws, baslik, aciklama, sutunlar):
        son_sutun = chr(64 + len(sutunlar))
        ws.sheet_view.showGridLines = False
        ws.merge_cells(f"A1:{son_sutun}1")
        ws["A1"] = baslik
        ws["A1"].fill = PatternFill("solid", fgColor="16324F")
        ws["A1"].font = Font(
            name="Aptos Display", color="FFFFFF", bold=True, size=17
        )
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 36
        ws.merge_cells(f"A2:{son_sutun}2")
        ws["A2"] = aciklama
        ws["A2"].font = Font(name="Aptos", color="64748B", size=10)
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        for sutun, metin in enumerate(sutunlar, start=1):
            hucre = ws.cell(4, sutun, metin)
            hucre.fill = PatternFill("solid", fgColor="2F6F77")
            hucre.font = Font(
                name="Aptos", color="FFFFFF", bold=True, size=10
            )
            hucre.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[chr(64 + sutun)].width = max(
                15, min(25, len(metin) + 4)
            )
        ws.row_dimensions[4].height = 27
        ws.freeze_panes = "A5"

    @classmethod
    def _rapor_satirlarini_yaz(cls, ws, satirlar):
        ince_cizgi = Side(style="thin", color="DDE5EC")
        for sira, satir in enumerate(satirlar, start=5):
            for sutun, deger in enumerate(satir, start=1):
                hucre = ws.cell(sira, sutun, deger)
                hucre.fill = PatternFill(
                    "solid", fgColor="F7FAFC" if sira % 2 else "FFFFFF"
                )
                hucre.font = Font(name="Aptos", color="334155", size=10)
                hucre.border = Border(bottom=ince_cizgi)
                hucre.alignment = Alignment(
                    horizontal="right" if isinstance(deger, (int, float)) else "left",
                    vertical="center",
                )
            ws.row_dimensions[sira].height = 22
        son_satir = max(4, 4 + len(satirlar))
        son_sutun = ws.max_column
        ws.auto_filter.ref = (
            f"A4:{chr(64 + son_sutun)}{son_satir}"
        )
        for sutun in range(1, son_sutun + 1):
            baslik = str(ws.cell(4, sutun).value or "")
            if "Hasılat" in baslik:
                for sira in range(5, son_satir + 1):
                    ws.cell(sira, sutun).number_format = '#,##0.00 "TL"'
            elif baslik in {"Toplam Araç", "Yaya Yolcu", "Araç", "Yaya"}:
                for sira in range(5, son_satir + 1):
                    ws.cell(sira, sutun).number_format = "#,##0"

    @classmethod
    def _grafik_sayfasini_hazirla(cls, ws, genel_satirlar, yon_satirlari):
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:J1")
        ws["A1"] = "FERİBOT RAPOR GRAFİKLERİ"
        ws["A1"].fill = PatternFill("solid", fgColor="16324F")
        ws["A1"].font = Font(
            name="Aptos Display", color="FFFFFF", bold=True, size=17
        )
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 36

        yon_toplamlari = defaultdict(lambda: [0, 0])
        for yon, _donem, arac, yaya, _hasilat in yon_satirlari:
            yon_toplamlari[yon][0] += arac
            yon_toplamlari[yon][1] += yaya
        ws.append([])
        ws.append([])
        ws.append(["Sefer Yönü", "Araç", "Yaya"])
        for yon in (MALATYA_BASKIL, BASKIL_MALATYA):
            arac, yaya = yon_toplamlari[yon]
            ws.append([yon, arac, yaya])

        ws["E4"] = "Dönem"
        ws["F4"] = "Araç"
        ws["G4"] = "Yaya"
        for sira, (donem, arac, yaya, _hasilat) in enumerate(
            reversed(genel_satirlar), start=5
        ):
            ws.cell(sira, 5, donem)
            ws.cell(sira, 6, arac)
            ws.cell(sira, 7, yaya)

        for hucre in list(ws[4])[:3] + list(ws[4])[4:7]:
            hucre.fill = PatternFill("solid", fgColor="2F6F77")
            hucre.font = Font(color="FFFFFF", bold=True)

        yon_grafigi = BarChart()
        yon_grafigi.title = "Yönlere Göre Toplam Geçiş"
        yon_grafigi.y_axis.title = "Adet"
        yon_grafigi.add_data(
            Reference(ws, min_col=2, max_col=3, min_row=4, max_row=6),
            titles_from_data=True,
        )
        yon_grafigi.set_categories(
            Reference(ws, min_col=1, min_row=5, max_row=6)
        )
        yon_grafigi.series[0].tx = SeriesLabel(v="Araç")
        yon_grafigi.series[1].tx = SeriesLabel(v="Yaya")
        yon_grafigi.height = 8
        yon_grafigi.width = 13
        ws.add_chart(yon_grafigi, "A9")

        if genel_satirlar:
            aylik_grafik = LineChart()
            aylik_grafik.title = "Aylık Genel Araç ve Yaya"
            aylik_grafik.y_axis.title = "Adet"
            son_satir = 4 + len(genel_satirlar)
            aylik_grafik.add_data(
                Reference(
                    ws, min_col=6, max_col=7,
                    min_row=4, max_row=son_satir
                ),
                titles_from_data=True,
            )
            aylik_grafik.set_categories(
                Reference(ws, min_col=5, min_row=5, max_row=son_satir)
            )
            aylik_grafik.series[0].tx = SeriesLabel(v="Araç")
            aylik_grafik.series[1].tx = SeriesLabel(v="Yaya")
            aylik_grafik.height = 8
            aylik_grafik.width = 13
            ws.add_chart(aylik_grafik, "H9")

        for sutun in ("A", "B", "C", "E", "F", "G"):
            ws.column_dimensions[sutun].width = 20

    @classmethod
    def eski_aylik_raporlari_temizle(cls, hedef_klasor=RAPOR_KLASORU):
        """Tek rapor merkezine geçildikten sonra üretilmiş eski aylıkları siler."""

        silinenler = []
        if not os.path.isdir(hedef_klasor):
            return silinenler
        for ad in os.listdir(hedef_klasor):
            if not (
                ad.startswith("Feribot_Aylik_Gecis_Raporu_")
                and ad.lower().endswith(".xlsx")
            ):
                continue
            dosya_yolu = os.path.join(hedef_klasor, ad)
            try:
                os.remove(dosya_yolu)
                silinenler.append(dosya_yolu)
            except OSError:
                # Excel'de açık olan eski dosya kapatıldığında sonraki
                # otomatik güncellemede yeniden temizlenir.
                continue
        return silinenler

    @classmethod
    def aylik_raporlari_guncelle(
        cls,
        gunluk_ozetler,
        hedef_klasor=RAPOR_KLASORU,
        referans_tarihi=None,
        sefer_yonu=MALATYA_BASKIL
    ):
        """Günlük toplamları aylarına göre ayrı Excel dosyalarında günceller."""

        aylik_ozetler = cls._ozetleri_aylara_ayir(gunluk_ozetler)

        if referans_tarihi is None:
            mevcut_tarihler = [
                tarih
                for satirlar in aylik_ozetler.values()
                for tarih, *_ in satirlar
            ]
            referans_tarihi = max(mevcut_tarihler, default=date.today())

        referans_degeri = cls._tarih_coz(referans_tarihi)
        if referans_degeri is None:
            raise ValueError("Rapor için geçerli bir referans tarihi gerekiyor.")

        aktif_ay = (referans_degeri.year, referans_degeri.month)
        aylik_ozetler.setdefault(aktif_ay, [])
        os.makedirs(hedef_klasor, exist_ok=True)

        guncellenen_dosyalar = []
        for (yil, ay), satirlar in sorted(aylik_ozetler.items()):
            dosya_yolu = os.path.join(
                hedef_klasor,
                cls.dosya_adi_olustur(yil, ay, sefer_yonu)
            )

            # Geçmiş aylar bir kez oluşturulur; yalnızca aktif ay sürekli güncellenir.
            if (yil, ay) != aktif_ay and os.path.exists(dosya_yolu):
                continue

            guncellenen_dosyalar.append(
                cls._aylik_raporu_yaz(
                    yil,
                    ay,
                    satirlar,
                    hedef_klasor,
                    sefer_yonu
                )
            )

        return guncellenen_dosyalar

    @classmethod
    def gunluk_raporu_guncelle(
        cls,
        gunluk_ozetler,
        hedef_klasor=RAPOR_KLASORU,
        referans_tarihi=None,
        sefer_yonu=MALATYA_BASKIL
    ):
        """Eski çağrılar için aktif aylık rapor dosyasının yolunu döndürür."""

        dosyalar = cls.aylik_raporlari_guncelle(
            gunluk_ozetler,
            hedef_klasor=hedef_klasor,
            referans_tarihi=referans_tarihi,
            sefer_yonu=sefer_yonu
        )
        return dosyalar[-1]

    @classmethod
    def aktar(cls, kayitlar, hedef_klasor=RAPOR_KLASORU):
        """Eski biçimdeki kayıtları ilgili aylık raporlara aktarır."""

        ozetler = cls.gunluk_ozet_hazirla(kayitlar)
        dosyalar = cls.aylik_raporlari_guncelle(
            ozetler,
            hedef_klasor=hedef_klasor
        )
        return dosyalar[-1]
